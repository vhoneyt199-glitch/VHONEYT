#!/usr/bin/env python3
"""Batch rename PDFs to Chinese titles with an Excel preview workflow.

Workflow:
  1. Generate rename_preview.xlsx:
       python scripts/rename_pdfs_chinese.py preview
  2. Open the Excel file, review/edit chinese_title, set confirm to Y.
  3. Apply renames:
       python scripts/rename_pdfs_chinese.py apply

Optional OpenAI translation during preview:
       set OPENAI_API_KEY=your_key
       python scripts/rename_pdfs_chinese.py preview --translator openai
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

PREVIEW_HEADERS = [
    "index",
    "confirm",
    "original_filename",
    "extracted_title",
    "chinese_title",
    "new_filename",
    "status",
]

INVALID_WINDOWS_CHARS = r'<>:"/\\|?*'
MAX_TITLE_CHARS = 120
DEFAULT_PREVIEW_NAME = "rename_preview.xlsx"


@dataclass
class PdfRenameRow:
    index: int
    confirm: str
    original_filename: str
    extracted_title: str
    chinese_title: str
    new_filename: str
    status: str


def default_pdf_folder() -> Path:
    """Return the user's Desktop/鼠标与键盘 folder on Windows-like machines."""
    userprofile = os.environ.get("USERPROFILE")
    if userprofile:
        return Path(userprofile) / "Desktop" / "鼠标与键盘"
    return Path.home() / "Desktop" / "鼠标与键盘"


def clean_cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_spaces(text: str) -> str:
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def sanitize_filename_part(text: str) -> str:
    text = normalize_spaces(text).replace("\n", " ")
    text = "".join("_" if char in INVALID_WINDOWS_CHARS else char for char in text)
    text = re.sub(r"\s+", " ", text).strip(" .")
    if len(text) > MAX_TITLE_CHARS:
        text = text[:MAX_TITLE_CHARS].rstrip(" .，,；;")
    return text or "未命名"


def has_chinese(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text))


def meaningful_lines(text: str) -> list[str]:
    lines: list[str] = []
    for raw_line in text.splitlines():
        line = normalize_spaces(raw_line)
        if not line:
            continue
        if len(line) < 4:
            continue
        if re.fullmatch(r"[\d\W_]+", line):
            continue
        lines.append(line)
    return lines


def looks_like_metadata_noise(text: str) -> bool:
    lowered = text.lower()
    noisy_words = [
        "doi:",
        "copyright",
        "abstract",
        "keywords",
        "journal",
        "vol.",
        "volume",
        "issue",
        "proceedings",
        "received",
        "accepted",
    ]
    return any(word in lowered for word in noisy_words)


def extract_title_from_first_page(pdf_path: Path) -> str:
    """Extract a likely paper title from metadata and first-page text."""
    from pypdf import PdfReader

    reader = PdfReader(str(pdf_path))

    metadata_title = ""
    if reader.metadata and reader.metadata.title:
        metadata_title = normalize_spaces(str(reader.metadata.title))
        if metadata_title and not looks_like_metadata_noise(metadata_title):
            return metadata_title

    if not reader.pages:
        return metadata_title

    first_page_text = normalize_spaces(reader.pages[0].extract_text() or "")
    lines = meaningful_lines(first_page_text)
    if not lines:
        return metadata_title

    # Most academic PDFs place the title in the first few meaningful lines.
    candidates: list[str] = []
    for start in range(min(8, len(lines))):
        candidate_parts: list[str] = []
        for line in lines[start : start + 4]:
            if looks_like_metadata_noise(line):
                break
            if re.search(r"@|\b(university|institute|department|school)\b", line, re.I):
                break
            candidate_parts.append(line)
            combined = " ".join(candidate_parts)
            if len(combined) >= 18:
                candidates.append(combined)
        if candidates:
            break

    if not candidates:
        candidates = lines[: min(3, len(lines))]

    candidate = max(candidates, key=len)
    candidate = re.sub(r"\s*-\s*ScienceDirect.*$", "", candidate, flags=re.I)
    return normalize_spaces(candidate)


def translate_with_openai(title: str, model: str) -> str:
    """Translate a title with OpenAI if the optional package and API key are available."""
    from openai import OpenAI

    client = OpenAI()
    response = client.responses.create(
        model=model,
        input=[
            {
                "role": "system",
                "content": (
                    "你是学术文献题名翻译助手。请把英文论文标题翻译为规范、简洁、自然的中文题名。"
                    "保留必要术语，不添加解释，不输出引号，不输出序号。"
                ),
            },
            {"role": "user", "content": title},
        ],
    )
    return normalize_spaces(response.output_text)


def propose_chinese_title(title: str, translator: str, openai_model: str) -> str:
    if not title:
        return ""
    if has_chinese(title):
        return title
    if translator == "openai":
        return translate_with_openai(title, openai_model)
    return ""


def build_new_filename(index: int, chinese_title: str) -> str:
    title = sanitize_filename_part(chinese_title) if chinese_title else "待填写中文题名"
    return f"{index:02d}-{title}.pdf"


def iter_pdf_files(folder: Path) -> Iterable[Path]:
    return sorted(
        (path for path in folder.iterdir() if path.is_file() and path.suffix.lower() == ".pdf"),
        key=lambda path: path.name.lower(),
    )


def write_preview(folder: Path, preview_path: Path, translator: str, openai_model: str) -> None:
    if not folder.exists():
        raise FileNotFoundError(f"PDF 文件夹不存在：{folder}")
    pdf_files = list(iter_pdf_files(folder))
    if not pdf_files:
        raise FileNotFoundError(f"该文件夹中没有 PDF 文件：{folder}")

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "rename_preview"
    worksheet.append(PREVIEW_HEADERS)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    for index, pdf_path in enumerate(pdf_files, start=1):
        status = "OK"
        extracted_title = ""
        chinese_title = ""
        try:
            extracted_title = extract_title_from_first_page(pdf_path)
            chinese_title = propose_chinese_title(extracted_title, translator, openai_model)
        except Exception as exc:
            status = f"需要人工检查：{exc}"
        new_filename = build_new_filename(index, chinese_title)
        worksheet.append(
            [
                index,
                "N",
                pdf_path.name,
                extracted_title,
                chinese_title,
                new_filename,
                status,
            ]
        )

    widths = [10, 10, 38, 70, 70, 70, 40]
    for offset, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + offset)].width = width
    worksheet.freeze_panes = "A2"
    workbook.save(preview_path)


def load_preview_rows(preview_path: Path) -> list[PdfRenameRow]:
    from openpyxl import load_workbook

    workbook = load_workbook(preview_path)
    worksheet = workbook.active
    headers = [clean_cell_text(cell.value) for cell in worksheet[1]]
    if headers[: len(PREVIEW_HEADERS)] != PREVIEW_HEADERS:
        raise ValueError(f"预览表表头不正确，应为：{PREVIEW_HEADERS}")

    rows: list[PdfRenameRow] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        values = list(row) + [""] * len(PREVIEW_HEADERS)
        rows.append(
            PdfRenameRow(
                index=int(values[0]),
                confirm=clean_cell_text(values[1]).upper(),
                original_filename=clean_cell_text(values[2]),
                extracted_title=clean_cell_text(values[3]),
                chinese_title=clean_cell_text(values[4]),
                new_filename=clean_cell_text(values[5]),
                status=clean_cell_text(values[6]),
            )
        )
    return rows


def unique_target_path(folder: Path, desired_name: str, source_path: Path) -> Path:
    desired_name = sanitize_filename_part(Path(desired_name).stem) + ".pdf"
    target = folder / desired_name
    if target.resolve() == source_path.resolve():
        return target
    if not target.exists():
        return target

    stem = target.stem
    suffix = target.suffix
    counter = 2
    while True:
        candidate = folder / f"{stem}（{counter}）{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def apply_renames(folder: Path, preview_path: Path, dry_run: bool) -> None:
    rows = load_preview_rows(preview_path)
    selected_rows = [row for row in rows if row.confirm in {"Y", "YES", "是", "1", "TRUE"}]
    if not selected_rows:
        print("没有需要重命名的行。请在 rename_preview.xlsx 中把 confirm 列改为 Y 后再执行。")
        return

    for row in selected_rows:
        if not row.chinese_title:
            print(f"跳过第 {row.index} 行：chinese_title 为空。")
            continue
        source_path = folder / row.original_filename
        if not source_path.exists():
            print(f"跳过第 {row.index} 行：原文件不存在：{source_path}")
            continue
        if row.new_filename and "待填写中文题名" not in row.new_filename:
            desired_name = row.new_filename
        else:
            desired_name = build_new_filename(row.index, row.chinese_title)
        target_path = unique_target_path(folder, desired_name, source_path)
        print(f"{source_path.name} -> {target_path.name}")
        if not dry_run:
            source_path.rename(target_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="将 PDF 文献按中文题名批量重命名。")
    subparsers = parser.add_subparsers(dest="command", required=True)

    def add_common_options(subparser: argparse.ArgumentParser) -> None:
        subparser.add_argument(
            "--folder",
            type=Path,
            default=default_pdf_folder(),
            help="PDF 文件夹路径，默认是当前用户桌面上的“鼠标与键盘”。",
        )
        subparser.add_argument(
            "--preview",
            type=Path,
            default=None,
            help="rename_preview.xlsx 路径；默认生成在 PDF 文件夹内。",
        )

    preview_parser = subparsers.add_parser("preview", help="生成 Excel 预览表，不重命名文件。")
    add_common_options(preview_parser)
    preview_parser.add_argument(
        "--translator",
        choices=["manual", "openai"],
        default="manual",
        help="翻译方式。manual 表示在 Excel 中人工填写；openai 表示使用 OPENAI_API_KEY 自动翻译。",
    )
    preview_parser.add_argument("--openai-model", default="gpt-4.1-mini", help="OpenAI 翻译模型。")

    apply_parser = subparsers.add_parser("apply", help="按 Excel 预览表执行重命名。")
    add_common_options(apply_parser)
    apply_parser.add_argument("--dry-run", action="store_true", help="只打印重命名计划，不真正改名。")

    return parser.parse_args()


def main() -> int:
    args = parse_args()
    folder = args.folder.expanduser().resolve()
    preview_path = args.preview.expanduser().resolve() if args.preview else folder / DEFAULT_PREVIEW_NAME

    if args.command == "preview":
        write_preview(folder, preview_path, args.translator, args.openai_model)
        print(f"已生成预览表：{preview_path}")
        print("请打开 Excel 检查 chinese_title/new_filename，并把确认重命名的行 confirm 改为 Y。")
        return 0

    if args.command == "apply":
        apply_renames(folder, preview_path, args.dry_run)
        return 0

    return 1


if __name__ == "__main__":
    sys.exit(main())
